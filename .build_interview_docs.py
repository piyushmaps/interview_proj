from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "interview-prep"
OUT.mkdir(parents=True, exist_ok=True)

# Status describes the repository inspected on October 1, 2026. Proposed means
# architecture we have discussed, not functionality already present in the code.
SECTIONS = [
    (
        "Product and system architecture",
        [
            ("What problem does this project solve, and who is it for?", "Both",
             "I am building an internal engineering knowledge assistant. An employee asks a question and gets a concise answer with links to relevant Jira issues, Slack discussions, and approved documents. The goal is to reduce time spent searching multiple systems. The first version is read-only: it retrieves and explains information, but does not create tickets or post messages."),
            ("Walk me through the current end-to-end architecture.", "Current code",
             "The current repository has two paths. During ingestion, source connectors produce normalized Documents; the pipeline hashes, chunks, embeds and writes them to FAISS or OpenSearch. During a question, FastAPI calls a LangGraph route-retrieve-synthesize graph. The router chooses indexed source labels, the retriever finds passages, and synthesis creates an answer with citations. Jira and Slack are currently part of the indexed-data design, not live MCP reads."),
            ("Why are ingestion and question answering separate workflows?", "Both",
             "Ingestion does expensive, repeatable preparation before employees ask questions: extract content, split it, embed it and publish it for search. A query then embeds a short question and searches prepared data. Separating the work reduces question latency and avoids reprocessing unchanged documents. The proposed change keeps this document pipeline while calling Jira and Slack live through MCP."),
            ("What is RAG, and where does this project use it?", "Current code",
             "Retrieval-augmented generation means finding source evidence first and giving that evidence to a language model to help it answer. Here, SourceRetriever returns relevant chunks, and synthesis receives those chunks with source names and URLs. The model is instructed to answer from the retrieved evidence. That instruction helps, but it does not by itself prove every claim is supported; citation and claim-support checks need strengthening."),
            ("Why does the project have connectors and a common Document model?", "Current code",
             "Jira, Slack, DevPortal and City Map have different shapes. Connectors translate them into one Pydantic Document with an ID, source, title, body, URL, author, update time and metadata. The rest of ingestion can process one contract instead of branching on every provider. That contract is for indexed documents; live MCP results need a shared evidence contract as a related but distinct model."),
            ("What is the difference between the vector index and the source of truth?", "Both",
             "The source system, such as Jira or Drive, is authoritative for its current record. A vector index contains derived passages and embeddings to make semantic search efficient; it can lag, omit content or have stale permissions. In the proposed design Jira and Slack are queried live, while document search uses an incrementally updated index. We should preserve source revision and retrieval time and recheck policy before serving indexed evidence."),
            ("What does the router do, and what does it not do?", "Current code",
             "The current router is a model call that returns source labels and optional subqueries. Code drops labels that are not in the SourceType enum, then retrieval filters the index by the selected labels. The router does not itself fetch a ticket or call a connector. In the proposed architecture, source specialists choose tools; our code authorizes each tool call and the MCP server executes it."),
            ("Why would this design use both live MCP retrieval and document indexing?", "Proposed design",
             "They solve different retrieval problems. Jira status and Slack discussions change frequently and are useful candidates for live search/read tools. Large reference documents are expensive to fetch and parse on every question, so we prepare an incremental searchable index for web pages, Drive files and office documents. The coordinator can use both paths for one question, then the evidence checker combines their results and limitations."),
            ("What are the most important boundaries between components?", "Both",
             "Source adapters translate provider data; policy code decides what we may use; retrieval finds candidate evidence; the answer model explains that evidence; and the API delivers a checked result. These boundaries make failures observable and tests focused. In the current code, the route-retrieve-synthesize boundary exists, but application authentication, eligibility enforcement, comprehensive sanitization and final output validation are not implemented."),
            ("What would you explain first to a new engineer joining this project?", "Both",
             "I would start with the two workflows: documents are prepared ahead of time, while employee questions retrieve evidence and generate an answer. Then I would trace one ticket-status question through the current FastAPI endpoint, LangGraph nodes and SourceRetriever. I would clearly label the MCP specialists, session memory and stronger guardrails as proposed design so the engineer does not mistake a diagram for shipped behavior."),
        ],
    ),
    (
        "Documents data and connectors",
        [
            ("What fields are in the current Document schema, and why?", "Current code",
             "Document has an ID, SourceType, title, body, URL, optional author and update timestamp, plus arbitrary metadata. The ID supports stable identity, source and title identify provenance, body is the indexable text, and URL allows citation. Metadata carries connector-specific details. For production I would define explicit revision, eligibility, redaction and processing-version fields rather than hiding security-critical state in arbitrary metadata."),
            ("Which source connectors exist, and how mature are they?", "Current code",
             "The repository contains Jira, Slack, DevPortal and City Map connectors with local sample data. The code provides a connector abstraction that yields Documents. That proves a local integration shape, not enterprise OAuth, permission synchronization, API quota compliance or production readiness. The new MCP exercise should use real protocol behavior against seeded fake data, then add provider-specific production authentication and policy review separately."),
            ("How would you add a new document source?", "Current code",
             "For the current ingestion design, I would implement iter_documents on a BaseConnector, map each source item into the Document contract, register the connector in build_connectors, and add fixtures and connector tests. For a production source I would also define stable IDs, revision/change detection, scoped credentials, pagination, deletion and permission events. The current interface alone does not supply those governance guarantees."),
            ("How should Google Drive and web pages enter the system?", "Proposed design",
             "I would create source adapters that enumerate approved scopes and revisions, fetch content under size and network limits, and emit a shared source-change record. Web fetching needs redirect and destination checks to prevent unintended network access. Drive requires scoped authentication and permission-change handling. Both flows then use the same extraction, eligibility, sanitization, versioning and publication stages."),
            ("How would you add Markdown, PDF and Word support?", "Both",
             "The current extractors include HTML, PDF, CSV or sheet-like data and email formats; there is no demonstrated Word parser in the inspected code. I would add format-specific extractors behind a common return type that includes text and location metadata. Markdown headings, PDF page numbers and Word heading structure should survive extraction for citations. Scanned PDFs need an explicitly enabled OCR step and separate accuracy evaluation."),
            ("Why does the project chunk documents, and how does it preserve context?", "Current code",
             "Chunking makes a long document searchable in smaller passages and bounds how much evidence enters a prompt. The configured default is 800 characters with 120 characters of overlap. The splitter prefers Markdown heading boundaries. Slack, Jira and City Map records are kept whole when short, because splitting a small conversation or row could remove context needed to interpret it."),
            ("Why is chunk identity deterministic?", "Current code",
             "Chunk IDs use the document ID plus a position, such as doc-42#3. Reprocessing a document can then overwrite matching chunk positions instead of appending duplicates. However, a document that shrinks can leave old trailing chunks unless all old chunks are deleted or a versioned publication scheme is used. The current pipeline deletes old chunks before writing replacements, which creates a failure window."),
            ("How does the project handle large spreadsheets?", "Current code",
             "City Map treats a table according to its content. Small reference sheets become row-level documents; large tables are profiled with schema, row count, column summaries and sample rows; large text-bearing sheets can also emit a bounded number of individual rows. This avoids embedding hundreds of thousands of nearly identical code or coordinate rows when a compact profile better matches likely questions."),
            ("How would you retain useful citation locations across extraction and chunking?", "Both",
             "I would attach structured location data to each text block and chunk, such as page number, Markdown heading, Word section or Drive item ID. The location should survive chunk creation and be included in the evidence contract. In the current Document schema, arbitrary metadata can hold details, but the chunking path does not yet define a complete normalized location contract for all formats."),
            ("How do deletion and permission changes differ from ordinary edits?", "Proposed design",
             "An edit builds and publishes a new version while the previous eligible version can remain available until the replacement is verified. A deletion or permission restriction must stop serving affected content promptly, so I would record a serving tombstone first, then invalidate caches and clean up indexed derivatives. A failed parse or partial source listing is not evidence that the source item was deleted."),
        ],
    ),
    (
        "Ingestion indexing and updates",
        [
            ("Walk me through one document through the current ingestion pipeline.", "Current code",
             "A connector yields a Document. The pipeline computes a SHA-256 hash of its indexable text, skips it if the manifest already has the same hash, otherwise chunks it. A batcher collects chunks, sends their text to the embedding service, initializes the vector store if needed, upserts chunks and records progress in the SQLite manifest. At the end it can prune records missing from the current enumeration and persists the store."),
            ("How does the manifest avoid embedding unchanged content?", "Current code",
             "The manifest is a SQLite table keyed by document ID, containing the last content hash, source, chunk count and indexed time. The pipeline hashes Document.to_text and compares it with the saved hash. An unchanged document is skipped before chunking and embedding. This makes repeat runs cheaper, although metadata-only changes such as changed eligibility are not part of that text hash and require stronger production version tracking."),
            ("How does the current pipeline recover from an interrupted run?", "Current code",
             "It commits manifest progress in batches, so a rerun can skip documents already recorded and continue processing changed or unseen documents. Embedding retries are bounded. The guarantee is not fully transactional: the manifest can be committed before a local FAISS persistence checkpoint. I would add crash-point tests and a durable version-publication protocol so metadata cannot claim a version is usable before its index data is durable."),
            ("What happens when an embedding request fails?", "Current code",
             "The _embed function retries up to four total attempts with exponential delays, then raises. That is intentional: silently reporting success for a partially indexed corpus is worse than stopping. A manifest and idempotent IDs make a later retry cheaper. Production would additionally classify retryable versus permanent errors, honor provider retry guidance, enforce the request/job deadline and record attempt budgets."),
            ("What is risky about how updated documents are written today?", "Current code",
             "For updates, _Batcher.flush deletes the old document chunks before upserting the new chunks. If the upsert fails, the previous searchable version is already gone. A safer design writes a candidate version first, verifies it, switches an active-version pointer, then removes the old version asynchronously. This needs reconciliation because metadata and vector storage do not share a transaction."),
            ("When is pruning safe, and what could go wrong?", "Current code",
             "Pruning compares the IDs seen in a run with IDs in the manifest, so it is safe only when the run completely enumerated the intended source scope. A partial API listing, permission error or parser failure must not be mistaken for deletion. The production connector contract should include scope identity and enumeration completeness, and tests should prove incomplete runs never prune."),
            ("Why use a streaming or generator-based ingestion pipeline?", "Current code",
             "Connectors and chunk generation can yield one document or passage at a time. This avoids retaining the full source corpus in memory; the embedding batch size and current document become the main memory controls. Streaming does not automatically solve durable job scheduling, backpressure or source checkpointing, so a larger system would put work behind a durable queue and independently scale workers."),
            ("Why batch embeddings, and what tradeoff does batch size create?", "Current code",
             "Batching amortizes request overhead and can improve throughput when sending multiple texts to the embedding service. A larger batch reduces the number of calls but uses more memory and means a failure may require retrying more work. The current default is 256 chunks. I would tune that using measured provider limits, latency, memory and retry cost rather than treating the default as a production fact."),
            ("How would you ingest changes continuously without re-embedding everything?", "Proposed design",
             "Use source change events or revision polling to enqueue idempotent jobs, compare source and processing revisions, and reprocess only changed items. Periodically reconcile each complete source scope in case events were missed. Track content, permissions, redaction policy, parser, chunker and embedding versions separately. Only a complete, verified new index version becomes active."),
            ("What ingestion metrics and tests would you add?", "Both",
             "I would measure source-to-search lag, discovered and processed changes, extraction and quarantine counts, embedding attempts, queue age, publication failures and deletion lag. Tests should cover duplicate and out-of-order changes, incomplete enumeration, corrupt files, permission revocation, model failure, and crashes before and after version activation. The current repository has unit and pipeline tests, but those production publication guarantees are still a design need."),
        ],
    ),
    (
        "Embeddings vector stores and retrieval",
        [
            ("What is an embedding, and where do query and document embeddings occur?", "Current code",
             "An embedding is a numeric vector representing text in a space where semantically similar text tends to be nearby. During ingestion, document chunks are embedded and stored with metadata. During a query, SourceRetriever embeds the question and asks the configured vector store to search. The question is embedded per lookup; documents are not re-embedded for every question."),
            ("How does cosine similarity work in this project?", "Current code",
             "Cosine similarity compares vector direction and is normalized between -1 and 1; larger scores mean closer directions. The VectorStore contract says every backend returns cosine-like relevance scores with higher being better. FAISS uses inner product over normalized vectors, which equals cosine similarity. The OpenSearch backend converts its reported score to the same convention so the retriever can apply one relevance policy."),
            ("What is top_k, and why is there also a similarity threshold?", "Current code",
             "top_k limits how many nearest passages are requested; its default is six. A vector index will return nearest neighbors even for a completely unrelated question, so the configured minimum similarity of 0.28 filters weak matches. The threshold supports an 'I could not find evidence' path. Both values are sample-corpus defaults and should be recalibrated using representative questions."),
            ("Why offer both FAISS and OpenSearch?", "Current code",
             "FAISS is the local default and requires no separate search server, making it convenient for learning and small corpora. The implementation keeps the index in memory and uses exact flat search. OpenSearch provides server-backed storage, native filtering and optional hybrid keyword-plus-vector search. Those are different scale and operational profiles; changing a backend alone does not prove capacity."),
            ("What limitation does FAISS metadata filtering create here?", "Current code",
             "The FAISS implementation cannot filter by source inside its nearest-neighbor search. It retrieves an expanding candidate window, filters candidates by source in application code and repeats until it finds enough matches or exhausts the index. That backend-specific workaround lives inside FaissVectorStore, leaving the retriever interface consistent. At large corpus sizes this can become expensive."),
            ("Why combine keyword and vector search?", "Current code",
             "Semantic search is useful for paraphrases, but may miss exact strings such as Jira keys, error codes or field names. Keyword search can recover those exact matches. OpenSearch supports an optional hybrid path that fuses rankings. The fusion and relevance thresholds must be evaluated on labeled queries; raw scores from different ranking methods should not be assumed to be directly comparable."),
            ("How does the retriever decide whether evidence is relevant?", "Current code",
             "It resolves the source filters and top_k, embeds the query, invokes the vector store and removes hits below min_similarity. That is intentionally small policy logic; FAISS-specific over-fetching belongs in its backend. A possible improvement is reranking a bounded candidate set using evaluated criteria, while keeping retrieval and answer generation separately measurable."),
            ("What are the current retrieval evaluation metrics and their limits?", "Current code",
             "The evaluation harness measures routing precision/recall and whether required golden facts appear in generated answers. The retrieval utility measures the first rank at which evidence containing each expected fact appears and calculates recall@k. It deliberately does not invent precision@k without passage relevance labels. The golden set is small, so these results do not represent enterprise traffic or corpus behavior."),
            ("What is a vector index not good at, and when would you use another retrieval method?", "Both",
             "Vector search is not an authority system, permission engine or guaranteed exact lookup. It can miss rare identifiers, return semantically similar but irrelevant text, and produce stale content. Use metadata filters and exact keyword search for IDs or constrained scopes; use relational queries for structured ownership/status facts when possible. Combine methods only where evaluation shows the added complexity improves recall or precision."),
            ("How would you scale retrieval from a laptop corpus to millions or billions of passages?", "Proposed design",
             "I would first measure passage count, bytes, vector dimensions, update rate, query rate, latency percentiles and filtering selectivity. Then benchmark compressed or disk/object-backed indexes and sharding/replication options under realistic workloads. I would isolate ingestion capacity from interactive reads, use quotas and admission controls, and validate freshness and access filters. A billion-passage design cannot be justified from the current FAISS sample."),
        ],
    ),
    (
        "LangGraph agents tools and orchestration",
        [
            ("What role does LangGraph play in the current project?", "Current code",
             "LangGraph wires three nodes: route, retrieve and synthesize. The same workflow is built synchronously for evaluation/scripts and asynchronously for the API. Nodes receive and update AgentState, and edges determine the next step. LangGraph controls workflow sequencing and streaming; it does not provide source authorization, memory policy or correctness automatically."),
            ("What is in the current AgentState?", "Current code",
             "AgentState is a TypedDict with question, selected_sources, sub_queries, retrieved chunks, answer, citations and routing_rationale. It has no session identifier, employee identity, checkpointer, permission context or request-wide budgets. Those belong in a deliberate future state design. Type annotations describe expected values but do not by themselves validate access policy or sanitize stored content."),
            ("Which current node is an agent, and which are ordinary steps?", "Current code",
             "The graph has a routing model call, a retrieval function and an answer model call. In a strict architecture sense, the current router is a bounded LLM task, not an autonomous tool-using specialist; retrieval is programmed code; synthesis is another bounded LLM task. The proposed redesign adds Jira and Slack specialists that can choose read tools in a limited loop. Those would be the tool-using agents."),
            ("How does the current router's output affect retrieval?", "Current code",
             "The router is prompted to return source names, subqueries and a rationale as JSON. _apply_routing drops unknown source names and falls back to DevPortal and City Map if none remain. Retrieval runs the subqueries against selected source labels and deduplicates returned chunks. The parse/fallback is not enough for production: query count, length, permissions and call budgets also need hard code validation."),
            ("Why separate an agent from its tools and policy gate?", "Proposed design",
             "The model can propose that it needs a search or record read; trusted application code should validate the server, tool, arguments, source scope and budget before calling MCP. The server executes only the approved call. This limits what prompt injection or a malformed plan can cause. A tool being visible through MCP does not automatically make it safe or permitted for every employee."),
            ("How would you structure bounded Jira and Slack specialists?", "Proposed design",
             "Each specialist gets a safe focused task and an explicit call/time budget. It proposes a read tool, code authorizes the exact call, the MCP client executes it, and the result is screened before it returns to the agent. The agent can stop when it has enough evidence or make a bounded follow-up. Tool results should not enter saved graph state before screening."),
            ("How would you run independent source agents in parallel and merge results?", "Proposed design",
             "After validating the plan, the graph can start selected Jira, Slack and document branches concurrently. Each branch should return its own typed result with evidence or a clear empty/failed status. A join waits within a deadline, combines updates by branch key and deduplicates by stable record/revision/location. Shared budgets must be atomic; parallel nodes must not overwrite stale copies of shared state."),
            ("What is conditional routing in LangGraph, and where would you use it?", "Both",
             "An edge chooses the next node based on current state. The current graph uses fixed route-to-retrieve-to-synthesize edges. The proposed graph would branch after the validated plan to selected source readers and after evidence assessment either compose an answer, refine once, or return a limitation. Every model-proposed branch still passes programmatic validation."),
            ("Would you give the answer-writing agent access to Jira and Slack tools?", "Proposed design",
             "I would not give the answer writer source tools. The specialists retrieve evidence under narrow, auditable permissions; the writer receives only approved evidence and produces a draft. That separation reduces unnecessary tool access and makes it easier to check where claims came from. If another retrieval is needed, it returns through the coordinator and the same gates."),
            ("How would you prevent an agent loop from running indefinitely?", "Proposed design",
             "Set a shared request deadline and aggregate limits for source calls, model calls, tokens, concurrency and refinement rounds. Validate before each call and charge dispatched attempts to a shared ledger that survives retries. Add a graph recursion limit as a backstop, typed failures and cancellation. A limit held only in one agent prompt is not a reliable control."),
        ],
    ),
    (
        "API streaming caching and session memory",
        [
            ("What API endpoints exist today?", "Current code",
             "The FastAPI application exposes POST /ask for a completed answer, POST /ask/stream for server-sent events, GET /health for basic status, and POST /cache/clear. The question request model checks a minimum length but not a maximum. The cache-clear endpoint is not protected by application authentication in the inspected code, so it must not be treated as production-secure."),
            ("Why are the API handlers asynchronous?", "Current code",
             "Question requests spend much of their time waiting for embedding, model and search services. The API uses async calls so an event loop can work on other requests during supported network waits. The retrieval method uses asyncio.to_thread for the synchronous vector-store search. Async improves concurrency for I/O waits; it does not increase provider quotas or CPU capacity without limit."),
            ("How does streaming work today?", "Current code",
             "The /ask/stream endpoint sends SSE events for started/status, synthesis tokens, completion and errors. astream_agent subscribes to LangGraph node updates and model messages, suppressing router JSON tokens and forwarding synthesis tokens. Citations arrive with the final result. A production concern is that draft text is already sent before a complete-output safety check, and the stream currently exposes exception text."),
            ("What are the risks of the current answer cache?", "Current code",
             "The cache is an in-process bounded LRU with a 15-minute default TTL and optional semantic matching using question embeddings. It stores the full answer and citations. Separate API instances do not share it, and its key does not include employee scope, policy version or evidence revisions. A stale or inappropriate cache hit could return an answer without rechecking current permissions."),
            ("How would you safely handle cache invalidation for live sources?", "Proposed design",
             "Track each answer's source IDs and revisions, policy and prompt versions, and freshness mode. On lookup, recheck current eligibility and dependency state; bypass caching for questions that require current live status unless the sources are revalidated. Start with caching disabled on the new retrieval path. A short TTL bounds age but does not immediately handle revocation."),
            ("What kinds of memory should this project distinguish?", "Proposed design",
             "Session memory is bounded safe context for follow-up references such as 'it means PAY-123.' LangGraph working state tracks the current request's tasks and evidence. Document knowledge is the searchable corpus. An answer cache is a performance optimization. None of these turns a previous generated answer into authoritative evidence; follow-up questions still retrieve fresh facts."),
            ("How would you isolate memory for thousands of employees?", "Proposed design",
             "Use a shared session store keyed by an opaque session ID and bound every read/write to the authenticated employee. Store limited sanitized context with an expiration time, and never trust an owner ID from the user or model. Enforce per-session update ordering, size limits and concurrency controls. Any LangGraph checkpoint must be associated with the same actor and rechecked on resume."),
            ("What happens if two requests update one session at once?", "Proposed design",
             "I would either serialize turns per session or use versioned writes so concurrent updates cannot silently overwrite each other. Each request reads a version and submits an update against that version; a conflict triggers a safe reload or asks the user to retry. Global requests from other employees remain concurrent. This behavior needs explicit tests before session memory is enabled."),
            ("How should failures appear during a stream?", "Both",
             "Before headers, return appropriate HTTP errors such as validation or service-unavailable responses. Once an SSE stream starts, send a terminal safe error event with a request ID, not an exception string. Distinguish a source timeout from a successful search with no results. Propagate client disconnects to cancel work where possible and release reservations."),
            ("How would you handle the five-second answer target?", "Both",
             "Treat it as a complete-answer objective, including citations, and define the percentile and load profile. Measure time in each graph stage, source/model calls, cache path and evidence count. Emit safe progress if work continues, but do not claim success from time-to-first-token alone. The current code has no demonstrated five-second percentile or concurrency benchmark."),
        ],
    ),
    (
        "Evaluation citations and observability",
        [
            ("What does the existing golden-set harness measure?", "Current code",
             "It compares selected sources against expected sources using recall and precision, and checks whether required facts appear in the answer after normalization. It can also compare with plain RAG and optionally judge groundedness. The harness helps diagnose routing versus generation. Its small sample fixtures do not prove production quality, security, or performance."),
            ("Why compare an agentic route with a plain RAG baseline?", "Current code",
             "The baseline asks retrieval for evidence across all sources and synthesizes an answer without routing. Comparing it with the routed graph helps determine whether the router improves evidence selection enough to justify its added model call, latency and failure modes. If routing does not improve measured quality or efficiency on representative evaluation questions, simplify it or improve its policy."),
            ("How would you evaluate document retrieval quality?", "Both",
             "Use questions with labeled relevant passages or facts and measure recall@k, reciprocal rank or nDCG, segmented by source and question type. Add exact identifier queries and paraphrases. The existing retrieval utility measures the rank of passages containing required golden facts, but it cannot compute passage precision without relevance labels. Maintain a held-out test set to avoid tuning on every example."),
            ("How would you evaluate live MCP search agents?", "Proposed design",
             "Create deterministic fake Jira and Slack servers with known records and realistic query behavior. Measure whether the right records are found, unnecessary calls made, pagination handled, and permission or error results interpreted correctly. Then contract-test against approved real servers. Keep source relevance and answer quality metrics separate so a fluent answer cannot hide missed retrieval."),
            ("How would you measure citation quality?", "Both",
             "Track citation precision—whether each cited source supports the nearby claim—and citation coverage—whether material factual claims have a source. Validate each citation ID and URL against evidence actually retrieved. Test missing, malformed and invented markers. The current _cited_only helper filters citation markers but does not prove claim support; if no markers appear, it returns all citations."),
            ("How would you test refusal and related-information behavior?", "Both",
             "Include off-topic questions, partial evidence, empty successful searches and unavailable sources as separate cases. The current synthesis path skips the answer-model call when no chunks clear the similarity cutoff. A richer design should distinguish 'no evidence found' from 'could not check source' and label related material without presenting it as proof. Measure both false answers and unnecessary refusals."),
            ("What does tracing currently capture, and what privacy concern follows?", "Current code",
             "Langfuse is optional and records model calls and retrieval traces when configured. The retrieval trace can include the query, sources, titles, URLs and a text preview. Those fields may contain sensitive information. Production telemetry needs data minimization, screening, access controls and retention rules; a tracing failure should not take down ordinary query handling."),
            ("What operational metrics would you add?", "Proposed design",
             "Measure completed-answer latency percentiles, safe time-to-first-event, errors and partial answers; per-source latency, 429s, timeouts and empty results; model tokens/cost; cache freshness/hits; ingestion lag and revocation lag; queue age; and citation/evaluation scores. Avoid raw questions or high-cardinality employee identifiers as metric labels. Pair alerts with an owner and a runbook."),
            ("How would you design a quality release gate?", "Proposed design",
             "Run a versioned golden set covering direct lookup, cross-source questions, conflicts, no-answer cases and injection fixtures. Report routing, retrieval recall, fact coverage, citation correctness, refusal behavior and latency separately. Set minimums only after a baseline is measured, review failures manually, and block a release for serious leakage or access-control regressions regardless of average answer score."),
            ("What is the difference between a unit test, a contract test and a load test here?", "Both",
             "A unit test checks a function such as citation filtering with fake inputs. A contract test checks our MCP adapter's protocol and expected tool schema against a local fake server or approved provider. A load test measures behavior under controlled concurrency and realistic downstream quotas. The current repository has substantial mocked unit tests; that is not evidence of 1,000-user capacity."),
        ],
    ),
    (
        "MCP proposed multi-agent design",
        [
            ("What does MCP provide in the proposed Jira and Slack integration?", "Proposed design",
             "MCP provides a standard client-server way for our application to discover and invoke tools and receive results. It does not automatically choose the right tool or write the final answer. We would use real local MCP servers backed by synthetic data first, then map the adapters to provider tool names and authentication. The server performs its tool; our application owns policy and answer composition."),
            ("Why use live Jira and Slack reads instead of only re-ingestion?", "Proposed design",
             "Current status and conversations change frequently, so live search/read can avoid depending on a periodically rebuilt full copy. It shifts challenges to network latency, quotas, permissions and failure handling. It also avoids re-embedding those live records as the query-time path. We still index documents because fetching and parsing large PDFs or Drive collections on every question would be costly and slow."),
            ("What fake MCP tools would you start with?", "Proposed design",
             "I would start with read-only search_issues and get_issue for Jira, and search_messages and get_thread for Slack. Seed linked tickets and threads, timestamps, missing fields, conflicting updates and empty results. Add pagination, delays, permission errors, malformed responses, rate limits and timeouts. The fake server should implement the actual protocol so the client lifecycle and tool calls are realistic."),
            ("How do you stop the model choosing a dangerous or unauthorized tool?", "Proposed design",
             "Do not give the model authority to execute arbitrary tools directly. Application code checks server and tool against an allowlist, validates exact arguments and scope, checks the actor's permissions, and reserves budget before the MCP call. Expose only read tools at first. Retrieved prompt-injection text cannot expand this allowlist or grant access."),
            ("What should a normalized evidence record contain across Jira, Slack and documents?", "Proposed design",
             "It should include a stable evidence ID, source and record ID, revision, sanitized title/text, canonical URL, location such as page or thread, source update time, retrieval time, and eligibility/redaction versions. Provider-specific details can remain metadata, but security and freshness fields should be explicit. The answer and citation renderer then use this common contract without discarding provenance."),
            ("How would you preserve freshness and handle conflicting sources?", "Proposed design",
             "Store source update and retrieval times separately. Present the newest relevant evidence first, but do not claim that newest means correct. If records conflict, describe each value with its date and link. If dates are missing or incomparable, say so. Current synthesis has a hard-coded preference for DevPortal over Slack and City Map for ownership, which should be reconciled with the freshness-first product requirement."),
            ("How many agents would you build, and why?", "Proposed design",
             "I would begin with two bounded tool-using specialists: Jira and Slack. Document search can remain deterministic retrieval code because its task is selecting passages from an index, not deciding among provider tools in a loop. Planning and answer writing are bounded model tasks. More agents add calls, latency and failure paths, so each should have a measured responsibility."),
            ("How do you combine successful evidence with partial source failures?", "Proposed design",
             "Each branch returns a typed outcome such as success, empty, skipped, denied, unavailable or timed out. A join combines evidence and source status separately. If Jira and documents worked but Slack timed out, answer the supported parts and clearly state Slack could not be checked. Never convert a timeout or permission denial into a claim that nothing exists."),
            ("What parallelism risks appear when Jira, Slack and document search run together?", "Proposed design",
             "Parallel branches reduce wall-clock time but multiply active provider calls and can race when writing shared graph state. I would cap per-request and global concurrency, atomically reserve shared budgets, assign each branch its own result key and merge only at a defined join. Cancellation and deadlines must stop or account for in-flight calls. Provider quotas are shared across application replicas."),
            ("When would you replace a live MCP search with an indexed copy or cache?", "Both",
             "If a source quota, latency or availability makes live reads unreliable, I would evaluate a compliant change-feed/index for data whose freshness and permission model can be maintained. A cache can help repeated questions but must track evidence revisions and access scope and recheck revocations. I would choose using measured latency, quotas, update frequency and policy requirements rather than avoiding ingestion categorically."),
        ],
    ),
    (
        "Security privacy and finance guardrails",
        [
            ("Is the project multi-user or multi-tenant?", "Proposed design",
             "The agreed product is one internal company service for many employees, so it is multi-user with per-employee access and session isolation. It is not currently designed as a SaaS service for multiple companies with separate organizations, credentials and policies. The current API code has no application authentication, so employee isolation is a required future feature, not something I would claim is already handled."),
            ("How should access control work when the assistant reads multiple sources?", "Proposed design",
             "Authenticate the employee at the API, then independently enforce that returned source content is eligible for all employees as required by the product. Source credentials and MCP tool capability do not themselves prove content is eligible. Check policy before model use, check current revocations before release, and bind session/cache entries to appropriate actor and policy scope."),
            ("What is prompt injection, and how does this architecture reduce its impact?", "Proposed design",
             "Prompt injection is text in a question or retrieved source that tries to redirect model behavior. Jira issues, Slack messages, PDFs and webpages are untrusted input. Screening helps but cannot guarantee detection, so the main controls are narrow read-only tools, code-enforced authorization and budgets, no credentials in prompts, bounded outputs, and checks before evidence reaches another model or employee."),
            ("How would you protect sensitive information in retrieved records?", "Both",
             "Check source eligibility and sensitivity separately. Screen titles, text, URLs and metadata before embedding, model prompts, checkpoints, traces or answer release. Redact only when the remaining content is still accurate; otherwise quarantine it. Use approved processing environments and avoid logging raw questions or records. Current code has no complete corporate classification integration, so proposed defaults require policy approval."),
            ("Where should the output guardrail sit, and why does streaming complicate it?", "Proposed design",
             "A complete-draft check should run before answer text is released. The current SSE path streams model tokens as they arrive, so later rejection cannot retract a sensitive token already delivered. The proposed flow streams safe progress first, checks the private complete draft for sensitive content and evidence support, then sends the checked answer. Sentence-level streaming would need its own safety design."),
            ("What data belongs in ordinary logs and traces?", "Proposed design",
             "Prefer request IDs, stage durations, counts, status codes, model versions and safe policy reason codes. Avoid full questions, document text, credentials and raw tool results. If security audit requires identifying who accessed what, use a restricted audit store with minimal stable actor/evidence references and controlled retention. The current tracing implementation can include text previews, so it needs a sanitization boundary."),
            ("How should secrets and provider credentials be managed?", "Both",
             "Keep credentials in a secret manager or local environment configuration for development; never put them in LangGraph state, prompts or returned citations. Use separate scoped service identities where possible, rotate credentials, restrict access and audit use. In production, tokens need correct issuer/audience validation; an MCP server should not accept arbitrary tokens and pass them through to another service."),
            ("What happens when permission checks or sensitive-data screening are unavailable?", "Proposed design",
             "Fail closed for the affected content: withhold it from the model and employee rather than assume it is safe. The response can explain that the source could not be checked without revealing restricted record details. Record a safe operational failure. For optional observability, degrade separately; telemetry failure should not cause raw data to be logged as a fallback."),
            ("How do you prevent stale or revoked evidence from reappearing through memory, cache or backups?", "Proposed design",
             "Maintain current revocation records keyed by stable source identity and apply them at retrieval and immediately before response release. Tie cached answers and session references to their evidence dependencies; invalidate affected entries when a record is restricted or deleted. Restores must reapply current tombstones before readiness. Physical backup purge follows approved policy and does not replace serving-time denial."),
            ("What security tests would be mandatory before a real finance-company pilot?", "Proposed design",
             "Test unauthorized identity/session access, restricted and unknown-eligibility records, synthetic secrets, prompt-injection content, unsafe URLs, malicious tool arguments, forged citations, cache leakage, deletion races and disabled policy services. Inspect every streamed event and trace payload, not only the final answer. Use approved synthetic data first, then require security and data-classification review before processing real corporate content."),
        ],
    ),
    (
        "Reliability deployment and growth",
        [
            ("What does the current health endpoint tell us, and what does it miss?", "Current code",
             "It reports whether a retriever was created and whether the startup chunk count was nonzero, along with vector-store, tracing and cache information. That is a useful local diagnostic but can become stale after startup and may expose internal details. Production needs separate liveness and current readiness checks, restricted diagnostics, safe dependency status and recovery behavior."),
            ("What does the Docker setup currently provide?", "Current code",
             "The Dockerfile builds a Python application image and Compose can run the API with a local OpenSearch service and mounted sample/index data. Compose explicitly disables the OpenSearch security plugin for local development. The Dockerfile does not declare a non-root user in the inspected file. These settings are convenient locally and should not be used as a production security configuration."),
            ("How would you make builds reproducible and safer?", "Proposed design",
             "Pin and lock transitive dependencies with hashes, build a minimal image, run as a non-root user, scan dependencies and image layers, and keep secrets out of image build arguments. Configure TLS and authentication in production and fail startup if development-only security flags are enabled. Promote the same tested artifact through stages with rollback support."),
            ("How would a load balancer and application instances support concurrency?", "Proposed design",
             "A load balancer routes each request to a healthy application instance; the instance runs the whole request graph. Add instances when measurements show query capacity is needed. Session state, budgets, rate limits and any shared cache must live in shared services or be consistently partitioned. A load balancer cannot increase Jira/Slack/model quotas, which remain shared downstream bottlenecks."),
            ("What does 1,000 concurrent users imply for capacity planning?", "Both",
             "First define whether 1,000 means signed-in people, active requests or requests arriving per second. If there are 1,000 in-flight requests and each takes five seconds on average, Little's Law suggests about 200 completions per second at steady state, but that is an illustration, not a benchmark. Multiple source/model calls per request increase downstream load. We must load-test realistic peaks and quotas."),
            ("How would you prevent overload from cascading across agents and providers?", "Proposed design",
             "Apply admission control before expensive work, bounded queues, per-user and global concurrency limits, shared provider quotas, request deadlines and cancellation. Retry only transient failures with bounded backoff and charge each dispatched attempt. Use circuit breakers to reduce traffic to an unhealthy source. Return a clear partial/unavailable result rather than letting queued work grow without limit."),
            ("How should you handle regional or index failure?", "Proposed design",
             "Keep a last known eligible index where policy allows, but apply current deletion and restriction state before serving it. Replicate metadata, job checkpoints, policy/tombstones, identity configuration and index artifacts according to recovery objectives. Measure restoration through drills and replay missed source changes. Replicating vectors alone is not a complete disaster-recovery plan."),
            ("What does the current AWS infrastructure documentation claim?", "Current code",
             "The repository has AWS architecture notes, but the handoff says Terraform/CDK deployment modules were not checked in. I would describe AWS as a direction under consideration, not a deployed system. Before final selection, verify current company approvals, service limits, cost, identity integration and recovery behavior, then add infrastructure-as-code and deployment tests."),
            ("How would you respond to a production incident where answers are slow?", "Proposed design",
             "Use traces and stage metrics to find whether API admission, session store, routing model, MCP source, document index or answer model is the bottleneck. Check saturation, timeouts, 429s and queue age; reduce optional fan-out or disable an unhealthy source using a documented switch. Keep the incident response safe: report missing sources, preserve audit, and avoid broad retries."),
            ("What is your phased roadmap from this repository to a production pilot?", "Both",
             "I would first establish an evaluation baseline and build a local real-protocol fake Jira MCP flow. Next add identity/policy boundaries and adversarial fixtures, then the Slack specialist and secure incremental document adapters. After that validate citations, memory isolation, cache invalidation, deletion/recovery, telemetry and deployment. Only then load-test progressively and qualify scale. Each phase has measurable exit criteria and review by data/security owners."),
        ],
    ),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, text, end):
        run._r.append(item)


def setup_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 45, 55)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)

    title_p = doc.add_paragraph(style="Title")
    title_p.paragraph_format.space_after = Pt(4)
    title_p.add_run(title)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run(subtitle)
    run.font.name = "Aptos"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(53, 95, 125)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header.add_run("INTUIT AGENTIC RAG  |  INTERVIEW PREPARATION")
    hrun.font.name = "Aptos"
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(90, 104, 116)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer.add_run("Project specific study guide  •  ")
    frun.font.size = Pt(8)
    frun.font.color.rgb = RGBColor(100, 110, 120)
    field(footer.add_run(), "PAGE")
    return doc


def intro_box(doc, heading, body):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.7)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF4F8")
    set_cell_borders(cell, "D6E1E8")
    set_cell_margins(cell, top=160, start=190, bottom=150, end=190)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(heading)
    run.bold = True
    run.font.color.rgb = RGBColor(25, 71, 101)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    return table


def add_question(doc, number, question, answer=None, status=None):
    h = doc.add_paragraph(style="Heading 2")
    h.paragraph_format.keep_with_next = answer is not None
    h.paragraph_format.keep_together = True
    h.add_run(f"Question {number:03d}")
    r = h.add_run()
    r.add_break()
    r = h.add_run(question)
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(35, 45, 55)
    if answer is not None:
        badge = doc.add_paragraph()
        badge.paragraph_format.space_after = Pt(2)
        badge.paragraph_format.keep_with_next = True
        br = badge.add_run(f"PROJECT STATUS  {status}")
        br.font.name = "Aptos"
        br.font.size = Pt(8)
        br.font.bold = True
        br.font.color.rgb = RGBColor(53, 95, 125)
        a = doc.add_paragraph()
        a.paragraph_format.left_indent = Inches(0.16)
        a.paragraph_format.space_after = Pt(9)
        a.paragraph_format.keep_together = True
        ar = a.add_run("Sample answer  ")
        ar.bold = True
        ar.font.color.rgb = RGBColor(25, 71, 101)
        a.add_run(answer)


def make_questions_doc():
    doc = setup_doc(
        "AI Engineering Interview Questions",
        "100 project specific questions for Intuit Agentic RAG",
    )
    intro_box(
        doc,
        "How to use this question bank",
        "Practice explaining the current code first, then discuss the proposed MCP and multi-agent design. The companion answer guide uses the same question numbers. Be explicit about what exists in the repository and what remains a design proposal.",
    )
    doc.add_paragraph(
        "The questions progress from product architecture to ingestion, retrieval, LangGraph, security and scale. Try answering each aloud before opening the matching answer."
    )
    number = 0
    for title, questions in SECTIONS:
        doc.add_paragraph(title, style="Heading 1")
        for question, _status, _answer in questions:
            number += 1
            add_question(doc, number, question)
    doc.core_properties.title = "AI Engineering Interview Questions"
    doc.core_properties.subject = "100 interview questions about the Intuit Agentic RAG project"
    doc.core_properties.keywords = "AI engineering, RAG, LangGraph, MCP, interview"
    doc.save(OUT / "AI_Engineering_Interview_Questions.docx")


def make_answers_doc():
    doc = setup_doc(
        "AI Engineering Interview Answers",
        "Model answers for the 100 project specific questions",
    )
    intro_box(
        doc,
        "Use these as speaking notes",
        "These are model answers, not claims that every planned capability is already built. The project status label distinguishes current code, proposed design and answers that discuss both. Put the idea in your own words and be ready to explain one concrete code path or tradeoff.",
    )
    doc.add_paragraph(
        "A strong interview answer states the design, gives evidence from the repository, and names a limitation or next step when one matters. Do not describe the fake MCP exercise or a design diagram as a production integration."
    )
    number = 0
    for title, questions in SECTIONS:
        doc.add_paragraph(title, style="Heading 1")
        for question, status, answer in questions:
            number += 1
            add_question(doc, number, question, answer=answer, status=status)
    doc.core_properties.title = "AI Engineering Interview Answers"
    doc.core_properties.subject = "Numbered sample answers about the Intuit Agentic RAG project"
    doc.core_properties.keywords = "AI engineering, RAG, LangGraph, MCP, interview, answers"
    doc.save(OUT / "AI_Engineering_Interview_Answers.docx")


if __name__ == "__main__":
    total = sum(len(items) for _title, items in SECTIONS)
    if total != 100:
        raise RuntimeError(f"Expected 100 question-answer pairs, found {total}")
    make_questions_doc()
    make_answers_doc()
    print(f"Created two documents with {total} aligned question-answer pairs in {OUT}")
